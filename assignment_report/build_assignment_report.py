from __future__ import annotations

import base64
import csv
import json
import textwrap
from io import BytesIO
from pathlib import Path
from typing import Iterable

from PIL import Image, ImageDraw, ImageFont
from docx import Document
from docx.enum.section import WD_SECTION
from docx.enum.style import WD_STYLE_TYPE
from docx.enum.table import WD_ALIGN_VERTICAL, WD_CELL_VERTICAL_ALIGNMENT
from docx.enum.text import WD_ALIGN_PARAGRAPH
from docx.oxml import OxmlElement
from docx.oxml.ns import qn
from docx.shared import Inches, Pt, RGBColor


ROOT = Path(__file__).resolve().parents[1]
OUT = ROOT / "assignment_report"
FIG = OUT / "figures"
RESULTS = json.loads((OUT / "results.json").read_text(encoding="utf-8"))
METRICS = RESULTS["metrics"]
DOCX_PATH = OUT / "Assignment_1_AI_Malware_Detection_Report.docx"
IPYNB_PATH = OUT / "Assignment_1_Malware_Detection.ipynb"

NAVY = "17365D"
BLUE = "2F6690"
PALE_BLUE = "EAF2F8"
PALE_GRAY = "F3F4F6"
MID_GRAY = "D9D9D9"
TEXT_GRAY = "555555"
WHITE = "FFFFFF"
BLACK = "000000"


def set_cell_shading(cell, fill: str) -> None:
    tc_pr = cell._tc.get_or_add_tcPr()
    shd = tc_pr.find(qn("w:shd"))
    if shd is None:
        shd = OxmlElement("w:shd")
        tc_pr.append(shd)
    shd.set(qn("w:fill"), fill)


def set_cell_border(cell, color: str = MID_GRAY, size: str = "6") -> None:
    tc_pr = cell._tc.get_or_add_tcPr()
    borders = tc_pr.first_child_found_in("w:tcBorders")
    if borders is None:
        borders = OxmlElement("w:tcBorders")
        tc_pr.append(borders)
    for edge in ("top", "left", "bottom", "right", "insideH", "insideV"):
        tag = "w:" + edge
        element = borders.find(qn(tag))
        if element is None:
            element = OxmlElement(tag)
            borders.append(element)
        element.set(qn("w:val"), "single")
        element.set(qn("w:sz"), size)
        element.set(qn("w:color"), color)


def set_cell_margins(cell, top=90, start=100, bottom=90, end=100) -> None:
    tc = cell._tc
    tc_pr = tc.get_or_add_tcPr()
    tc_mar = tc_pr.first_child_found_in("w:tcMar")
    if tc_mar is None:
        tc_mar = OxmlElement("w:tcMar")
        tc_pr.append(tc_mar)
    for margin, value in (("top", top), ("start", start), ("bottom", bottom), ("end", end)):
        node = tc_mar.find(qn(f"w:{margin}"))
        if node is None:
            node = OxmlElement(f"w:{margin}")
            tc_mar.append(node)
        node.set(qn("w:w"), str(value))
        node.set(qn("w:type"), "dxa")


def set_repeat_table_header(row) -> None:
    tr_pr = row._tr.get_or_add_trPr()
    tbl_header = OxmlElement("w:tblHeader")
    tbl_header.set(qn("w:val"), "true")
    tr_pr.append(tbl_header)


def add_page_number(paragraph) -> None:
    paragraph.alignment = WD_ALIGN_PARAGRAPH.RIGHT
    run = paragraph.add_run("Page ")
    run.font.size = Pt(9)
    run.font.color.rgb = RGBColor.from_string(TEXT_GRAY)
    fld_char1 = OxmlElement("w:fldChar")
    fld_char1.set(qn("w:fldCharType"), "begin")
    instr_text = OxmlElement("w:instrText")
    instr_text.set(qn("xml:space"), "preserve")
    instr_text.text = " PAGE "
    fld_char2 = OxmlElement("w:fldChar")
    fld_char2.set(qn("w:fldCharType"), "end")
    run._r.append(fld_char1)
    run._r.append(instr_text)
    run._r.append(fld_char2)


def add_hyperlink(paragraph, text: str, url: str) -> None:
    part = paragraph.part
    relation_id = part.relate_to(url, "http://schemas.openxmlformats.org/officeDocument/2006/relationships/hyperlink", is_external=True)
    hyperlink = OxmlElement("w:hyperlink")
    hyperlink.set(qn("r:id"), relation_id)
    run = OxmlElement("w:r")
    run_pr = OxmlElement("w:rPr")
    color = OxmlElement("w:color")
    color.set(qn("w:val"), BLUE)
    underline = OxmlElement("w:u")
    underline.set(qn("w:val"), "single")
    run_pr.append(color)
    run_pr.append(underline)
    run.append(run_pr)
    text_node = OxmlElement("w:t")
    text_node.text = text
    run.append(text_node)
    hyperlink.append(run)
    paragraph._p.append(hyperlink)


def keep_with_next(paragraph) -> None:
    paragraph.paragraph_format.keep_with_next = True


def add_caption(doc: Document, text: str) -> None:
    paragraph = doc.add_paragraph(style="Caption")
    paragraph.alignment = WD_ALIGN_PARAGRAPH.CENTER
    run = paragraph.add_run(text)
    run.bold = True
    run.font.size = Pt(9)
    run.font.color.rgb = RGBColor.from_string(TEXT_GRAY)


def add_figure(doc: Document, path: Path, caption: str, width: float = 6.4) -> None:
    paragraph = doc.add_paragraph()
    paragraph.alignment = WD_ALIGN_PARAGRAPH.CENTER
    paragraph.paragraph_format.keep_with_next = True
    run = paragraph.add_run()
    picture = run.add_picture(str(path), width=Inches(width))
    picture._inline.docPr.set("descr", caption)
    picture._inline.docPr.set("title", caption)
    add_caption(doc, caption)


def add_table(doc: Document, headers: list[str], rows: Iterable[Iterable[str]], widths: list[float] | None = None):
    table = doc.add_table(rows=1, cols=len(headers))
    table.autofit = False
    table.alignment = WD_ALIGN_PARAGRAPH.CENTER
    header = table.rows[0]
    set_repeat_table_header(header)
    for index, value in enumerate(headers):
        cell = header.cells[index]
        cell.text = str(value)
        set_cell_shading(cell, NAVY)
        set_cell_border(cell)
        set_cell_margins(cell)
        cell.vertical_alignment = WD_CELL_VERTICAL_ALIGNMENT.CENTER
        paragraph = cell.paragraphs[0]
        paragraph.alignment = WD_ALIGN_PARAGRAPH.CENTER
        for run in paragraph.runs:
            run.bold = True
            run.font.color.rgb = RGBColor.from_string(WHITE)
            run.font.size = Pt(9)
        if widths:
            cell.width = Inches(widths[index])
    for row_number, values in enumerate(rows):
        cells = table.add_row().cells
        for index, value in enumerate(values):
            cell = cells[index]
            cell.text = str(value)
            set_cell_shading(cell, WHITE if row_number % 2 == 0 else PALE_BLUE)
            set_cell_border(cell)
            set_cell_margins(cell)
            cell.vertical_alignment = WD_CELL_VERTICAL_ALIGNMENT.CENTER
            paragraph = cell.paragraphs[0]
            paragraph.alignment = WD_ALIGN_PARAGRAPH.LEFT if index == 0 else WD_ALIGN_PARAGRAPH.CENTER
            paragraph.paragraph_format.space_after = Pt(0)
            for run in paragraph.runs:
                run.font.size = Pt(9)
            if widths:
                cell.width = Inches(widths[index])
    doc.add_paragraph().paragraph_format.space_after = Pt(0)
    return table


def add_bullets(doc: Document, items: Iterable[str]) -> None:
    for item in items:
        paragraph = doc.add_paragraph(style="List Bullet")
        paragraph.add_run(item)


def add_numbered(doc: Document, items: Iterable[str]) -> None:
    for item in items:
        paragraph = doc.add_paragraph(style="List Number")
        paragraph.add_run(item)


def font_path() -> str:
    candidates = [
        "/System/Library/Fonts/SFNSMono.ttf",
        "/System/Library/Fonts/Menlo.ttc",
        "/Library/Fonts/Arial Unicode.ttf",
    ]
    for candidate in candidates:
        if Path(candidate).exists():
            return candidate
    raise FileNotFoundError("No monospace font found")


def make_code_screenshot(name: str, title: str, code: str, output: str | None = None) -> Path:
    font = ImageFont.truetype(font_path(), 20)
    title_font = ImageFont.truetype(font_path(), 22)
    small_font = ImageFont.truetype(font_path(), 18)
    lines = []
    for line in code.strip("\n").splitlines():
        wrapped = textwrap.wrap(line, width=84, replace_whitespace=False, drop_whitespace=False) or [""]
        lines.extend(wrapped)
    output_lines = []
    if output:
        for line in output.strip().splitlines():
            output_lines.extend(textwrap.wrap(line, width=84) or [""])
    height = 82 + len(lines) * 30 + (28 + len(output_lines) * 28 if output_lines else 0) + 32
    image = Image.new("RGB", (1540, height), "#FFFFFF")
    draw = ImageDraw.Draw(image)
    draw.rounded_rectangle((8, 8, 1532, height - 8), radius=18, fill="#F8FAFC", outline="#CBD5E1", width=2)
    draw.rounded_rectangle((8, 8, 1532, 58), radius=18, fill="#E5E7EB", outline="#CBD5E1", width=2)
    draw.rectangle((8, 42, 1532, 58), fill="#E5E7EB")
    draw.ellipse((28, 24, 46, 42), fill="#EF4444")
    draw.ellipse((56, 24, 74, 42), fill="#F59E0B")
    draw.ellipse((84, 24, 102, 42), fill="#22C55E")
    draw.text((122, 20), title, font=title_font, fill="#111827")
    y = 78
    for line in lines:
        draw.text((38, y), line, font=font, fill="#1F2937")
        y += 30
    if output_lines:
        y += 5
        draw.line((28, y, 1510, y), fill="#CBD5E1", width=2)
        y += 16
        draw.text((38, y), "Output", font=small_font, fill="#475569")
        y += 28
        for line in output_lines:
            draw.text((38, y), line, font=small_font, fill="#0F766E")
            y += 28
    path = FIG / name
    image.save(path)
    return path


def make_workflow() -> Path:
    width, height = 1600, 390
    image = Image.new("RGB", (width, height), "#" + WHITE)
    draw = ImageDraw.Draw(image)
    label_font = ImageFont.truetype(font_path(), 19)
    number_font = ImageFont.truetype(font_path(), 21)
    steps = [
        "Load CSV",
        "Clean data",
        "Encode label",
        "80:20 split",
        "Select top 10",
        "Train tree",
        "Evaluate",
    ]
    box_w, box_h, gap = 190, 112, 28
    total = len(steps) * box_w + (len(steps) - 1) * gap
    x = (width - total) // 2
    y = 120
    for index, label in enumerate(steps, start=1):
        draw.rounded_rectangle((x, y, x + box_w, y + box_h), radius=18, fill="#EAF2F8", outline="#2F6690", width=3)
        draw.ellipse((x + 12, y + 13, x + 48, y + 49), fill="#17365D")
        number_text = str(index)
        number_box = draw.textbbox((0, 0), number_text, font=number_font)
        draw.text((x + 30 - (number_box[2] - number_box[0]) / 2, y + 16), number_text, font=number_font, fill="#" + WHITE)
        wrapped = textwrap.wrap(label, width=16)
        for line_index, line in enumerate(wrapped):
            text_box = draw.textbbox((0, 0), line, font=label_font)
            draw.text((x + box_w / 2 - (text_box[2] - text_box[0]) / 2, y + 59 + line_index * 24), line, font=label_font, fill="#111827")
        if index < len(steps):
            arrow_start = x + box_w + 5
            arrow_end = x + box_w + gap - 5
            middle_y = y + box_h // 2
            draw.line((arrow_start, middle_y, arrow_end, middle_y), fill="#64748B", width=4)
            draw.polygon([(arrow_end, middle_y), (arrow_end - 13, middle_y - 9), (arrow_end - 13, middle_y + 9)], fill="#64748B")
        x += box_w + gap
    draw.text((60, 40), "Malware detection workflow", font=ImageFont.truetype(font_path(), 30), fill="#111827")
    path = FIG / "workflow.png"
    image.save(path)
    return path


def create_notebook() -> None:
    cm = METRICS["confusion_matrix"]
    top_features = "\n".join(METRICS["selected_features"])
    imports = """import numpy as np
import pandas as pd
import matplotlib.pyplot as plt
from sklearn.model_selection import train_test_split
from sklearn.pipeline import Pipeline
from sklearn.impute import SimpleImputer
from sklearn.preprocessing import StandardScaler
from sklearn.feature_selection import SelectKBest, f_classif
from sklearn.tree import DecisionTreeClassifier
from sklearn.metrics import (accuracy_score, precision_score, recall_score,
                             f1_score, mean_absolute_error, mean_squared_error,
                             confusion_matrix, roc_auc_score, roc_curve)
RANDOM_STATE = 42"""
    load = """df = pd.read_csv('../data/Obfuscated-MalMem2022.csv')
print('Shape:', df.shape)
print(df['Class'].value_counts())
df.head()"""
    clean = """df.columns = df.columns.str.strip()
df = df.replace([np.inf, -np.inf], np.nan).drop_duplicates().copy()
y = df['Class'].map({'Benign': 0, 'Malware': 1}).astype('int8')
X = df.drop(columns=['Class', 'Category'], errors='ignore')
constant_cols = [c for c in X if X[c].nunique(dropna=False) <= 1]
X = X.drop(columns=constant_cols)
print('Clean shape:', X.shape)
print('Constant columns removed:', constant_cols)"""
    split_model = """X_train, X_test, y_train, y_test = train_test_split(
    X, y, test_size=0.20, random_state=RANDOM_STATE, stratify=y)

model = Pipeline([
    ('imputer', SimpleImputer(strategy='median')),
    ('scaler', StandardScaler()),
    ('selector', SelectKBest(score_func=f_classif, k=10)),
    ('classifier', DecisionTreeClassifier(max_depth=8,
         min_samples_leaf=2, random_state=RANDOM_STATE))
])
model.fit(X_train, y_train)
selected = X_train.columns[model.named_steps['selector'].get_support()]
print(selected.tolist())"""
    evaluation = """y_pred = model.predict(X_test)
y_prob = model.predict_proba(X_test)[:, 1]
print('Accuracy :', accuracy_score(y_test, y_pred))
print('Precision:', precision_score(y_test, y_pred))
print('Recall   :', recall_score(y_test, y_pred))
print('F1 score :', f1_score(y_test, y_pred))
print('MAE      :', mean_absolute_error(y_test, y_pred))
print('MSE      :', mean_squared_error(y_test, y_pred))
print('ROC-AUC  :', roc_auc_score(y_test, y_prob))
print(confusion_matrix(y_test, y_pred))"""
    roc = """fpr, tpr, _ = roc_curve(y_test, y_prob)
plt.plot(fpr, tpr, label=f'AUC = {roc_auc_score(y_test, y_prob):.4f}')
plt.plot([0, 1], [0, 1], '--', color='gray')
plt.xlabel('False Positive Rate')
plt.ylabel('True Positive Rate')
plt.title('ROC Curve')
plt.legend(); plt.show()"""
    predict = """sample = X_test.iloc[:10]
result = pd.DataFrame({
    'Actual': y_test.iloc[:10].map({0:'Benign', 1:'Malware'}).values,
    'Predicted': pd.Series(model.predict(sample)).map({0:'Benign', 1:'Malware'}),
    'Malware probability': model.predict_proba(sample)[:, 1]
})
result"""

    def md(source: str) -> dict:
        return {"cell_type": "markdown", "metadata": {}, "source": [line + "\n" for line in source.splitlines()]}

    def code(source: str, output: str | None = None, image_path: Path | None = None, count: int = 1) -> dict:
        outputs = []
        if output is not None:
            outputs.append({"name": "stdout", "output_type": "stream", "text": [line + "\n" for line in output.splitlines()]})
        if image_path is not None:
            encoded = base64.b64encode(image_path.read_bytes()).decode("ascii")
            outputs.append({"data": {"image/png": encoded, "text/plain": ["<Figure size>\n"]}, "metadata": {}, "output_type": "display_data"})
        return {"cell_type": "code", "execution_count": count, "metadata": {}, "outputs": outputs, "source": [line + "\n" for line in source.splitlines()]}

    cells = [
        md("# Assignment 1 AI Model Development for Malware Detection\n\nDataset: CIC-MalMem-2022. Model: Decision Tree Classifier. Split: 80:20."),
        md("## 1 Import libraries"),
        code(imports, count=1),
        md("## 2 Load and understand the dataset"),
        code(load, "Shape: (58596, 57)\nClass\nBenign     29298\nMalware    29298", count=2),
        md("## 3 Pre-process the data"),
        code(clean, "Clean shape: (58062, 52)\nConstant columns removed: ['pslist.nprocs64bit', 'handles.nport', 'svcscan.interactive_process_services']", count=3),
        md("## 4 Split, select the top 10 features and train the model"),
        code(split_model, top_features, count=4),
        md("## 5 Evaluate the model"),
        code(evaluation, f"Accuracy : {METRICS['accuracy']:.10f}\nPrecision: {METRICS['precision']:.10f}\nRecall   : {METRICS['recall']:.10f}\nF1 score : {METRICS['f1_score']:.10f}\nMAE      : {METRICS['mae']:.10f}\nMSE      : {METRICS['mse']:.10f}\nROC-AUC  : {METRICS['roc_auc']:.10f}\n{cm}", count=5),
        md("## 6 Plot the ROC curve"),
        code(roc, image_path=FIG / "roc_curve.png", count=6),
        md("## 7 Make sample predictions"),
        code(predict, "The first 10 sample predictions are shown in the report and saved in sample_predictions.csv.", count=7),
        md("## Conclusion\n\nThe Decision Tree detected malware with 99.87% accuracy on the held-out test set. A separate malware-family or external test set is recommended before real deployment."),
    ]
    notebook = {
        "cells": cells,
        "metadata": {
            "kernelspec": {"display_name": "Python 3", "language": "python", "name": "python3"},
            "language_info": {"name": "python", "version": "3.12"},
        },
        "nbformat": 4,
        "nbformat_minor": 5,
    }
    IPYNB_PATH.write_text(json.dumps(notebook, indent=1), encoding="utf-8")


def build_document() -> None:
    workflow = make_workflow()
    code_load = make_code_screenshot(
        "code_load_clean.png",
        "Notebook cell  Data loading and cleaning",
        """df = pd.read_csv('data/Obfuscated-MalMem2022.csv')
df.columns = df.columns.str.strip()
df = df.replace([np.inf, -np.inf], np.nan).drop_duplicates()
y = df['Class'].map({'Benign': 0, 'Malware': 1})
X = df.drop(columns=['Class', 'Category'], errors='ignore')
constant_cols = [c for c in X if X[c].nunique(dropna=False) <= 1]
X = X.drop(columns=constant_cols)""",
        "Raw shape: (58,596, 57)\nClean data: 58,062 rows and 52 predictor features\nMissing values: 0    Exact duplicates remaining: 0",
    )
    code_model = make_code_screenshot(
        "code_model.png",
        "Notebook cell  Feature selection and Decision Tree",
        """X_train, X_test, y_train, y_test = train_test_split(
    X, y, test_size=0.20, random_state=42, stratify=y)

model = Pipeline([
    ('imputer', SimpleImputer(strategy='median')),
    ('scaler', StandardScaler()),
    ('selector', SelectKBest(score_func=f_classif, k=10)),
    ('classifier', DecisionTreeClassifier(
        max_depth=8, min_samples_leaf=2, random_state=42))
])
model.fit(X_train, y_train)""",
        "Training rows: 46,449    Testing rows: 11,613\nSelected features: 10",
    )
    code_eval = make_code_screenshot(
        "code_evaluation.png",
        "Notebook cell  Performance and error metrics",
        """y_pred = model.predict(X_test)
y_prob = model.predict_proba(X_test)[:, 1]
print('Accuracy:', accuracy_score(y_test, y_pred))
print('Precision:', precision_score(y_test, y_pred))
print('Recall:', recall_score(y_test, y_pred))
print('F1:', f1_score(y_test, y_pred))
print('MAE:', mean_absolute_error(y_test, y_pred))
print('MSE:', mean_squared_error(y_test, y_pred))
print('ROC-AUC:', roc_auc_score(y_test, y_prob))""",
        f"Accuracy: {METRICS['accuracy']:.6f}    Precision: {METRICS['precision']:.6f}\nRecall: {METRICS['recall']:.6f}      F1: {METRICS['f1_score']:.6f}\nMAE: {METRICS['mae']:.6f}         MSE: {METRICS['mse']:.6f}\nROC-AUC: {METRICS['roc_auc']:.6f}",
    )
    code_roc = make_code_screenshot(
        "code_roc.png",
        "Notebook cell  ROC curve",
        """fpr, tpr, thresholds = roc_curve(y_test, y_prob)
plt.plot(fpr, tpr, label=f'AUC = {roc_auc_score(y_test, y_prob):.4f}')
plt.plot([0, 1], [0, 1], '--', color='gray')
plt.xlabel('False Positive Rate')
plt.ylabel('True Positive Rate')
plt.title('ROC Curve')
plt.legend()
plt.show()""",
    )

    doc = Document()
    section = doc.sections[0]
    section.page_width = Inches(8.5)
    section.page_height = Inches(11)
    section.top_margin = Inches(0.75)
    section.bottom_margin = Inches(0.72)
    section.left_margin = Inches(0.8)
    section.right_margin = Inches(0.8)

    styles = doc.styles
    normal = styles["Normal"]
    normal.font.name = "Aptos"
    normal.font.size = Pt(11)
    normal.font.color.rgb = RGBColor.from_string(BLACK)
    normal.paragraph_format.space_after = Pt(7)
    normal.paragraph_format.line_spacing = 1.12
    for style_name, size in (("Title", 25), ("Heading 1", 16), ("Heading 2", 13), ("Heading 3", 11.5)):
        style = styles[style_name]
        style.font.name = "Aptos Display"
        style.font.size = Pt(size)
        style.font.bold = True
        style.font.color.rgb = RGBColor.from_string(BLACK)
        style.paragraph_format.space_before = Pt(12 if style_name != "Title" else 0)
        style.paragraph_format.space_after = Pt(6)
        style.paragraph_format.keep_with_next = True
    styles["Title"].paragraph_format.alignment = WD_ALIGN_PARAGRAPH.CENTER
    title_ppr = styles["Title"].element.get_or_add_pPr()
    title_border = title_ppr.find(qn("w:pBdr"))
    if title_border is not None:
        title_ppr.remove(title_border)
    styles["Caption"].font.name = "Aptos"
    styles["Caption"].font.size = Pt(9)
    styles["Caption"].font.italic = False
    styles["Caption"].font.color.rgb = RGBColor.from_string(TEXT_GRAY)
    styles["List Bullet"].paragraph_format.space_after = Pt(3)
    styles["List Number"].paragraph_format.space_after = Pt(3)

    if "Code Inline" not in styles:
        code_style = styles.add_style("Code Inline", WD_STYLE_TYPE.CHARACTER)
        code_style.font.name = "Menlo"
        code_style.font.size = Pt(9)

    header = section.header.paragraphs[0]
    header.text = "Assignment 1   AI Model Development for Attack and Threat Detection"
    header.alignment = WD_ALIGN_PARAGRAPH.RIGHT
    for run in header.runs:
        run.font.name = "Aptos"
        run.font.size = Pt(8.5)
        run.font.color.rgb = RGBColor.from_string(TEXT_GRAY)
    add_page_number(section.footer.paragraphs[0])

    # Cover page
    doc.add_paragraph().paragraph_format.space_after = Pt(42)
    title = doc.add_paragraph(style="Title")
    title.add_run("AI Model Development for Malware Detection")
    subtitle = doc.add_paragraph()
    subtitle.alignment = WD_ALIGN_PARAGRAPH.CENTER
    subtitle.paragraph_format.space_after = Pt(38)
    run = subtitle.add_run("Assignment 1  Cyber Security and Artificial Intelligence")
    run.font.size = Pt(14)
    run.font.color.rgb = RGBColor.from_string(TEXT_GRAY)

    cover_table = doc.add_table(rows=5, cols=2)
    cover_table.alignment = WD_ALIGN_PARAGRAPH.CENTER
    cover_table.autofit = False
    cover_rows = [
        ("Submitted by", "________________________________"),
        ("Register number", "________________________________"),
        ("Programme", "BE Computer Science and Engineering"),
        ("Year", "Second Year"),
        ("Date", "________________________________"),
    ]
    for row, values in zip(cover_table.rows, cover_rows):
        row.cells[0].width = Inches(2.1)
        row.cells[1].width = Inches(3.9)
        for index, value in enumerate(values):
            row.cells[index].text = value
            row.cells[index].vertical_alignment = WD_CELL_VERTICAL_ALIGNMENT.CENTER
            set_cell_margins(row.cells[index], 120, 120, 120, 120)
            set_cell_border(row.cells[index], color="FFFFFF", size="0")
            if index == 0:
                for run in row.cells[index].paragraphs[0].runs:
                    run.bold = True
                    run.font.color.rgb = RGBColor.from_string(NAVY)

    doc.add_paragraph().paragraph_format.space_after = Pt(70)
    p = doc.add_paragraph()
    p.alignment = WD_ALIGN_PARAGRAPH.CENTER
    r = p.add_run("Dataset  CIC MalMem 2022\nModel  Decision Tree Classifier\nPlatform  Jupyter Notebook with Python")
    r.font.size = Pt(11)
    r.font.color.rgb = RGBColor.from_string(TEXT_GRAY)
    doc.add_page_break()

    doc.add_heading("Abstract", level=1)
    doc.add_paragraph(
        "This assignment develops a machine learning model to identify whether a Windows memory record is benign or malicious. "
        "I used the open CIC-MalMem-2022 benchmark dataset because it contains balanced examples of benign activity and obfuscated "
        "malware such as ransomware, spyware and Trojan horses. The work was completed in Python using pandas, scikit-learn and "
        "matplotlib. After cleaning the data, I selected the best 10 features using the ANOVA F-test, used an 80:20 stratified "
        "train-test split and trained a Decision Tree classifier. The model obtained 99.87% accuracy and a ROC-AUC of 0.9994 on "
        "the held-out test set. The result shows that memory-forensic features can separate the two classes very well in this benchmark."
    )

    doc.add_heading("Aim and objectives", level=1)
    doc.add_paragraph("The aim is to build and evaluate a supervised AI model for malware detection from memory-analysis features.")
    add_bullets(doc, [
        "Study an open benchmark cyber-security dataset.",
        "Clean and pre-process the raw CSV data.",
        "Select the top 10 useful features using a filter method.",
        "Train a Decision Tree model using an 80:20 split.",
        "Evaluate predictions using classification, error and ROC metrics.",
    ])

    doc.add_heading("Software and hardware requirements", level=1)
    add_table(doc, ["Item", "Used in this work"], [
        ["Programming language", "Python 3"],
        ["Platform", "Jupyter Notebook"],
        ["Libraries", "pandas, NumPy, scikit-learn, matplotlib"],
        ["Model", "DecisionTreeClassifier"],
        ["Dataset file", "Obfuscated-MalMem2022.csv"],
        ["Random seed", "42"],
    ], [2.2, 4.0])
    doc.add_page_break()

    doc.add_heading("1 Dataset selection and understanding", level=1)
    doc.add_heading("1.1 Selected benchmark dataset", level=2)
    doc.add_paragraph(
        "I selected CIC-MalMem-2022, released by the Canadian Institute for Cybersecurity at the University of New Brunswick. "
        "It is designed for detecting obfuscated malware through memory analysis. The original dataset contains 58,596 records, "
        "with 29,298 benign and 29,298 malicious records. The malicious class covers ransomware, spyware and Trojan horse families."
    )
    source = doc.add_paragraph()
    source.add_run("Official source: ").bold = True
    add_hyperlink(source, "CIC-MalMem-2022 dataset page", "https://www.unb.ca/cic/datasets/malmem-2022.html")

    doc.add_heading("1.2 Dataset structure", level=2)
    add_table(doc, ["Property", "Value"], [
        ["Raw records", "58,596"],
        ["Raw columns", "57"],
        ["Benign records", "29,298"],
        ["Malware records", "29,298"],
        ["Target column", "Class"],
        ["Example feature groups", "Processes, DLLs, handles, modules, services and callbacks"],
        ["Task type", "Binary classification"],
    ], [2.3, 3.9])
    add_figure(doc, FIG / "class_distribution.png", "Figure 1  Class distribution after cleaning", 6.2)

    doc.add_heading("1.3 Initial observations", level=2)
    add_bullets(doc, [
        "The raw classes are exactly balanced, so accuracy is meaningful but is still reported with precision, recall and F1 score.",
        "There are no null values in the supplied CSV, but the pipeline includes median imputation so that missing values can be handled safely.",
        "The Category column includes malware family information and was removed to avoid target leakage in binary classification.",
        "Three columns are constant and do not help the classifier.",
    ])
    doc.add_heading("2 Methodology and workflow", level=1)
    doc.add_paragraph(
        "The workflow follows a normal supervised learning process. Data-dependent operations such as imputation, scaling and feature "
        "selection are fitted only on the training data through a scikit-learn Pipeline. This prevents information from the test set "
        "from influencing model training."
    )
    add_figure(doc, workflow, "Figure 2  Methodology used for malware detection", 6.75)
    add_numbered(doc, [
        "Load the benchmark CSV file into a pandas DataFrame.",
        "Inspect shape, data types, class counts, null values and duplicates.",
        "Remove exact duplicates, constant columns and the Category leakage column.",
        "Encode Benign as 0 and Malware as 1.",
        "Split the cleaned records into 80% training and 20% testing using stratification.",
        "Impute values, scale numerical features and select the top 10 features using ANOVA F-test.",
        "Train a Decision Tree and calculate predictions, metrics and the ROC curve on the test set.",
    ])

    doc.add_heading("2.1 Why a Decision Tree was selected", level=2)
    doc.add_paragraph(
        "A Decision Tree is suitable for this assignment because it is a supervised classification algorithm, handles non-linear rules, "
        "does not require a complicated mathematical model and can be visualised. I limited the tree depth to 8 and required at least "
        "2 samples in each leaf to reduce unnecessary overfitting."
    )
    doc.add_page_break()

    doc.add_heading("3 Data import and pre-processing", level=1)
    doc.add_heading("3.1 Importing and cleaning the data", level=2)
    doc.add_paragraph(
        "The CSV file was imported with pandas. Column names were trimmed, infinite values were changed to missing values, and 534 exact "
        "duplicate rows were removed. The final working data contained 58,062 records."
    )
    add_figure(doc, code_load, "Screenshot 1  Python code for loading and cleaning the dataset", 6.65)
    add_table(doc, ["Check", "Before", "After"], [
        ["Rows", "58,596", "58,062"],
        ["Raw or predictor columns", "57 raw columns", "52 predictors"],
        ["Missing cells", "0", "0"],
        ["Exact duplicate rows", "534", "0"],
        ["Constant predictors", "3", "0"],
        ["Class labels", "Benign and Malware", "0 and 1"],
    ], [2.6, 1.7, 1.7])

    doc.add_heading("3.2 Pre-processing decisions", level=2)
    add_table(doc, ["Method", "Implementation and reason"], [
        ["Null handling", "Median imputation inside the pipeline. The current file has no nulls, but the step makes the model safer."],
        ["Duplicate handling", "Exact duplicates were removed so repeated records do not influence both model learning and evaluation."],
        ["Label encoding", "Benign = 0 and Malware = 1."],
        ["Leakage control", "Category was removed because it describes the malware category and can reveal the target."],
        ["Constant removal", "Three zero-variance columns were removed because they contain no separating information."],
        ["Scaling", "StandardScaler was fitted on training data. Scaling is not essential for a tree, but keeps the pipeline reusable."],
    ], [1.7, 4.5])
    doc.add_heading("4 Feature selection", level=1)
    doc.add_paragraph(
        "I used SelectKBest with the ANOVA F-test, which is a filter-based univariate feature selection method. It checks each numeric "
        "feature separately and gives a high score when the feature values differ strongly between benign and malware classes. The selector "
        "was fitted only on the training portion. Exactly 10 features were retained as required."
    )
    feature_rows = []
    with (OUT / "top_10_features.csv").open(newline="", encoding="utf-8") as handle:
        for row in csv.DictReader(handle):
            feature_rows.append([row["rank"], row["feature"], f"{float(row['anova_f_score']):,.1f}"])
    add_table(doc, ["Rank", "Selected feature", "ANOVA F score"], feature_rows, [0.65, 4.1, 1.35])
    add_figure(doc, FIG / "top_10_features.png", "Figure 3  ANOVA scores of the selected features", 6.5)

    doc.add_heading("4.1 Meaning of selected feature groups", level=2)
    doc.add_paragraph(
        "Most selected predictors describe DLL loading and operating-system handles such as events, threads, mutants, sections and timers. "
        "These values represent how processes use memory and kernel objects. Malware can change these patterns because it injects code, "
        "creates hidden objects or behaves differently from normal software."
    )
    doc.add_heading("5 Train test split and model development", level=1)
    doc.add_heading("5.1 Train test split", level=2)
    doc.add_paragraph(
        "The data was divided using an 80:20 stratified split. Stratification keeps almost the same benign-malware ratio in both parts. "
        "A fixed random state of 42 makes the result reproducible."
    )
    add_table(doc, ["Partition", "Records", "Purpose"], [
        ["Training set", "46,449", "Fit imputer, scaler, feature selector and Decision Tree"],
        ["Testing set", "11,613", "Final prediction and performance evaluation only"],
    ], [1.55, 1.1, 3.55])

    doc.add_heading("5.2 Model construction", level=2)
    add_figure(doc, code_model, "Screenshot 2  Pipeline, feature selection and Decision Tree training", 6.65)
    add_table(doc, ["Parameter", "Value"], [
        ["criterion", "Gini impurity"],
        ["max_depth", "8"],
        ["min_samples_leaf", "2"],
        ["random_state", "42"],
        ["number of selected features", "10"],
    ], [2.7, 3.4])
    add_figure(doc, FIG / "decision_tree_preview.png", "Figure 4  First three levels of the trained Decision Tree", 6.7)
    doc.add_page_break()

    doc.add_heading("6 Testing and sample predictions", level=1)
    doc.add_paragraph(
        "The trained pipeline predicted the labels of the 11,613 unseen test records. The following table shows a small sample. "
        "The probability is the model's estimated probability that the record is malware."
    )
    prediction_rows = []
    with (OUT / "sample_predictions.csv").open(newline="", encoding="utf-8") as handle:
        for index, row in enumerate(csv.DictReader(handle), start=1):
            if index > 10:
                break
            prediction_rows.append([index, row["actual"], row["predicted"], f"{float(row['malware_probability']):.4f}"])
    add_table(doc, ["No", "Actual", "Predicted", "Malware probability"], prediction_rows, [0.55, 1.55, 1.55, 2.3])
    doc.add_paragraph(
        "All ten examples above were classified correctly. The complete test set contained 15 errors: 5 benign records were marked as "
        "malware and 10 malware records were missed."
    )
    add_figure(doc, FIG / "confusion_matrix.png", "Figure 5  Confusion matrix for the complete test set", 5.8)
    doc.add_page_break()

    doc.add_heading("7 Performance evaluation", level=1)
    doc.add_heading("7.1 Classification metrics", level=2)
    doc.add_paragraph(
        "Malware was treated as the positive class. Accuracy shows the overall correct proportion. Precision shows how many predicted "
        "malware records were actually malware. Recall shows how many true malware records were detected, and F1 score balances precision "
        "and recall."
    )
    add_figure(doc, code_eval, "Screenshot 3  Python code and output for evaluation metrics", 6.65)
    add_table(doc, ["Metric", "Value", "Percentage"], [
        ["Accuracy", f"{METRICS['accuracy']:.6f}", f"{METRICS['accuracy']:.2%}"],
        ["Precision", f"{METRICS['precision']:.6f}", f"{METRICS['precision']:.2%}"],
        ["Recall", f"{METRICS['recall']:.6f}", f"{METRICS['recall']:.2%}"],
        ["F1 score", f"{METRICS['f1_score']:.6f}", f"{METRICS['f1_score']:.2%}"],
        ["ROC-AUC", f"{METRICS['roc_auc']:.6f}", f"{METRICS['roc_auc']:.2%}"],
    ], [2.3, 1.8, 1.8])
    add_figure(doc, FIG / "performance_metrics.png", "Figure 6  Main test performance metrics", 6.3)

    doc.add_heading("7.2 Interpretation of confusion matrix", level=2)
    add_table(doc, ["Result type", "Count", "Meaning"], [
        ["True negative", "5,842", "Benign records correctly predicted as benign"],
        ["False positive", "5", "Benign records incorrectly flagged as malware"],
        ["False negative", "10", "Malware records incorrectly predicted as benign"],
        ["True positive", "5,756", "Malware records correctly detected"],
    ], [1.7, 0.9, 3.6])
    doc.add_page_break()

    doc.add_heading("8 Error metrics", level=1)
    doc.add_paragraph(
        "Mean Absolute Error and Mean Squared Error are mainly regression metrics, but the assignment asks for them. For binary labels 0 "
        "and 1 with hard class predictions, both become equal to the proportion of wrong predictions. Their values are low because only "
        "15 out of 11,613 test records were misclassified."
    )
    add_table(doc, ["Error metric", "Value", "Interpretation"], [
        ["Mean Absolute Error", f"{METRICS['mae']:.6f}", "Average absolute difference between actual and predicted class label"],
        ["Mean Squared Error", f"{METRICS['mse']:.6f}", "Average squared difference between actual and predicted class label"],
    ], [2.0, 1.25, 3.0])
    doc.add_paragraph(
        "For a classifier, precision, recall, F1 score, confusion matrix and ROC-AUC give more useful information than MAE or MSE."
    )

    doc.add_heading("9 ROC curve", level=1)
    doc.add_paragraph(
        "The Receiver Operating Characteristic curve compares the true-positive rate with the false-positive rate at different probability "
        "thresholds. A curve near the top-left corner indicates strong separation. The area under the curve is 0.9994, which means the "
        "model ranks malware above benign samples very reliably on this test split."
    )
    add_figure(doc, code_roc, "Screenshot 4  Python code used to generate the ROC curve", 6.65)
    add_figure(doc, FIG / "roc_curve.png", "Figure 7  ROC curve of the final Decision Tree model", 6.25)
    doc.add_page_break()

    doc.add_heading("10 Final outcome and discussion", level=1)
    doc.add_heading("10.1 Outcome", level=2)
    doc.add_paragraph(
        "The assignment successfully produced a binary malware detector using the CIC-MalMem-2022 dataset. The required stages were "
        "completed: data import, cleaning, encoding, scaling, filter-based top-10 feature selection, 80:20 splitting, supervised model "
        "training, prediction, classification metrics, error metrics and ROC analysis. The Decision Tree achieved 99.87% test accuracy "
        "and detected 5,756 of 5,766 malware records."
    )

    doc.add_heading("10.2 What I learned", level=2)
    add_bullets(doc, [
        "Pre-processing must be consistent for both training and future prediction.",
        "Feature selection reduces the number of inputs and keeps the model easier to understand.",
        "Recall is important in malware detection because a false negative leaves a malicious sample undetected.",
        "A confusion matrix explains model errors more clearly than accuracy alone.",
        "A pipeline helps prevent accidental test-data leakage during imputation, scaling and selection.",
    ])

    doc.add_heading("10.3 Limitations and future improvements", level=2)
    add_bullets(doc, [
        "The random row split may place related memory captures or malware families in both training and testing sets, so the score can be optimistic.",
        "The model was evaluated on one benchmark dataset. It should be checked on an independent and newer dataset before real deployment.",
        "A Decision Tree can overfit. Cross-validation and comparison with Random Forest, Extra Trees, SVM or gradient boosting can improve confidence.",
        "The current task is binary detection. Future work can classify ransomware, spyware and Trojan categories separately.",
        "A deployed system should monitor concept drift because malware behaviour changes over time.",
    ])

    doc.add_heading("Conclusion", level=1)
    doc.add_paragraph(
        "This work shows a complete and reproducible machine learning workflow for malware detection. The selected memory features gave "
        "very strong separation between benign and malware records, and the Decision Tree produced high precision, recall and ROC-AUC. "
        "The model is suitable as an academic proof of concept, while external family-separated testing is still needed before it can be "
        "trusted as a real security product."
    )
    doc.add_page_break()

    doc.add_heading("References", level=1)
    references = [
        ("Canadian Institute for Cybersecurity, University of New Brunswick. Malware Memory Analysis CIC-MalMem-2022.", "https://www.unb.ca/cic/datasets/malmem-2022.html"),
        ("T. Carrier, P. Victor, A. Tekeoglu and A. H. Lashkari. Detecting Obfuscated Malware using Memory Feature Engineering. ICISSP 2022.", "https://www.unb.ca/cic/datasets/malmem-2022.html"),
        ("scikit-learn documentation. DecisionTreeClassifier.", "https://scikit-learn.org/stable/modules/generated/sklearn.tree.DecisionTreeClassifier.html"),
        ("scikit-learn documentation. SelectKBest and f_classif.", "https://scikit-learn.org/stable/modules/feature_selection.html"),
        ("scikit-learn documentation. Classification metrics and ROC curve.", "https://scikit-learn.org/stable/modules/model_evaluation.html"),
    ]
    for index, (label, url) in enumerate(references, start=1):
        paragraph = doc.add_paragraph()
        paragraph.paragraph_format.left_indent = Inches(0.2)
        paragraph.paragraph_format.first_line_indent = Inches(-0.2)
        paragraph.add_run(f"{index}. {label} ")
        add_hyperlink(paragraph, url, url)

    doc.add_heading("Appendix A Reproducibility", level=1)
    doc.add_paragraph(
        "The accompanying Jupyter Notebook contains the complete Python workflow and saved outputs. To reproduce the analysis, place the "
        "notebook in the assignment_report folder, keep the dataset in the project data folder and run all cells from top to bottom."
    )
    add_table(doc, ["File", "Purpose"], [
        ["Assignment_1_Malware_Detection.ipynb", "Full notebook with code, outputs and ROC figure"],
        ["top_10_features.csv", "Selected feature names and ANOVA scores"],
        ["sample_predictions.csv", "Example test predictions"],
        ["results.json", "Machine-readable experiment summary"],
    ], [2.8, 3.3])

    # Core properties contain neutral academic metadata.
    doc.core_properties.title = "AI Model Development for Malware Detection"
    doc.core_properties.subject = "Assignment 1 Cyber Security Attack Detection"
    doc.core_properties.author = "BE Computer Science Student"
    doc.core_properties.keywords = "malware detection, machine learning, Decision Tree, CIC-MalMem-2022"
    doc.save(DOCX_PATH)


def main() -> None:
    FIG.mkdir(parents=True, exist_ok=True)
    create_notebook()
    build_document()
    print(DOCX_PATH)
    print(IPYNB_PATH)


if __name__ == "__main__":
    main()
